#!/usr/bin/env python3
"""Build the expanded Yuanyuan Hu CV as DOCX and PDF from cv-data.json."""
import html
import json
import os
import textwrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate


HERE = os.path.dirname(os.path.abspath(__file__))
DATA_PATH = os.path.join(HERE, "cv-data.json")
DOCX_OUT = os.path.join(HERE, "Yuanyuan_Hu_CV_Extended.docx")
PDF_OUT = os.path.join(HERE, "Yuanyuan_Hu_CV_Extended.pdf")

NAVY = RGBColor(0x1E, 0x3A, 0x5F)
SOFT_BLUE = RGBColor(0x3A, 0x5E, 0x8C)
INK = RGBColor(0x26, 0x2A, 0x2E)
MUTED = RGBColor(0x5F, 0x66, 0x6E)
BODY_FONT = "Calibri"
HEAD_FONT = "Cambria"


def norm_url(url):
    return (url or "").replace("https://", "").rstrip("/") + "/" if url else ""


def doi_url(pub):
    doi = (pub.get("links") or {}).get("doi", "")
    if not doi:
        return ""
    return doi if doi.startswith("http") else f"https://doi.org/{doi}"


def set_run_font(run, name=BODY_FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)


def add_run(p, text, *, bold=False, italic=False, size=9.5, color=INK, font=BODY_FONT):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.color.rgb = color
    set_run_font(r, font)
    return r


def tight(p, before=0, after=4, line=1.08):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def section_title(doc, title):
    p = doc.add_paragraph()
    tight(p, before=8, after=3)
    add_run(p, title.upper(), bold=True, size=10.5, color=NAVY, font=HEAD_FONT)
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C9D3E0")
    pbdr.append(bottom)
    ppr.append(pbdr)
    return p


def build_docx(data):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.58)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(footer_p, f"Latest CV: {norm_url(data['site']['url'])}", size=7.2, color=MUTED)

    doc.styles["Normal"].font.name = BODY_FONT
    doc.styles["Normal"].font.size = Pt(9.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(p, after=0)
    add_run(p, data["name"], bold=True, size=20, color=NAVY, font=HEAD_FONT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(p, after=0)
    add_run(p, data["title"], italic=True, size=10.5, color=SOFT_BLUE, font=HEAD_FONT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(p, after=1)
    contacts = [
        data.get("email", ""),
        data.get("orcid", "").replace("https://", ""),
        data.get("researchgate", "").replace("https://www.", "").replace("https://", ""),
        data.get("location", ""),
    ]
    add_run(p, "  |  ".join(c for c in contacts if c), size=8.6, color=SOFT_BLUE)

    section_title(doc, "Research Profile")
    p = doc.add_paragraph()
    tight(p, line=1.12)
    add_run(p, data.get("profile", ""), size=9.4, color=INK)

    section_title(doc, "Research Positioning")
    p = doc.add_paragraph()
    tight(p, after=2)
    add_run(p, "Recommended primary positioning: ", bold=True, size=9.2, color=INK)
    add_run(p, data["positioningOptions"][0], size=9.2, color=INK)
    for option in data.get("positioningOptions", [])[1:]:
        p = doc.add_paragraph(style=None)
        tight(p, after=1, line=1.05)
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        add_run(p, "- ", size=9, color=SOFT_BLUE)
        add_run(p, option, size=9, color=INK)

    section_title(doc, "Education")
    for e in data["education"]:
        p = doc.add_paragraph()
        tight(p, after=0)
        add_run(p, e["degree"], bold=True, size=9.4, color=INK)
        add_run(p, f"    {e['period']}", size=8.8, color=SOFT_BLUE)
        p = doc.add_paragraph()
        tight(p, after=0)
        add_run(p, e["institution"], size=9, color=INK)
        if e.get("location"):
            add_run(p, f", {e['location']}", size=9, color=MUTED)
        p = doc.add_paragraph()
        tight(p, after=3, line=1.08)
        add_run(p, e.get("detail", ""), size=8.8, color=MUTED)
        if e.get("supervisor"):
            add_run(p, f" Supervisor: {e['supervisor']}.", size=8.8, color=MUTED)

    section_title(doc, "Research Interests and Methods")
    p = doc.add_paragraph()
    tight(p, after=2, line=1.08)
    add_run(p, "Interests: ", bold=True, size=9.2, color=INK)
    add_run(p, "; ".join(data.get("interests", [])) + ".", size=9.2, color=INK)
    p = doc.add_paragraph()
    tight(p, after=2, line=1.08)
    add_run(p, "Methods: ", bold=True, size=9.2, color=INK)
    add_run(p, " ".join(data.get("skills", [])), size=9.2, color=INK)

    section_title(doc, "Awards and Honours")
    for award in data.get("awards", []):
        p = doc.add_paragraph()
        tight(p, after=1, line=1.06)
        p.paragraph_format.left_indent = Inches(0.45)
        p.paragraph_format.first_line_indent = Inches(-0.45)
        add_run(p, f"{award['year']}  ", bold=True, size=8.8, color=SOFT_BLUE)
        add_run(p, award["text"], size=8.8, color=INK)

    section_title(doc, "Publications")
    pubs = data.get("publications", {})
    for label, key in (("Lead-author publications", "lead"), ("Collaborative publications", "collaborative")):
        if not pubs.get(key):
            continue
        p = doc.add_paragraph()
        tight(p, before=2, after=2)
        add_run(p, label, bold=True, size=9, color=MUTED)
        for i, pub in enumerate(pubs[key], 1):
            p = doc.add_paragraph()
            tight(p, after=2, line=1.08)
            p.paragraph_format.left_indent = Inches(0.24)
            p.paragraph_format.first_line_indent = Inches(-0.24)
            add_run(p, f"{i}. ", bold=True, size=8.5, color=SOFT_BLUE)
            add_run(p, f"{pub['authors']} ({pub['year']}). {pub['title']}. ", size=8.5, color=INK)
            add_run(p, pub["venue"], italic=True, size=8.5, color=INK)
            add_run(p, f", {pub.get('info', '')}. ", size=8.5, color=INK)
            details = [pub.get("contribution"), pub.get("metrics")]
            add_run(p, " ".join(f"[{d}]" for d in details if d), size=7.7, color=MUTED)
            if doi_url(pub):
                p = doc.add_paragraph()
                tight(p, after=2, line=1.0)
                p.paragraph_format.left_indent = Inches(0.24)
                add_run(p, "DOI/article: ", bold=True, size=7.5, color=MUTED)
                add_run(p, doi_url(pub), size=7.5, color=SOFT_BLUE)

    doc.save(DOCX_OUT)


def para(text, style):
    return Paragraph(text, style)


def safe(text):
    return html.escape(text or "")


def highlighted_authors(authors):
    return safe(authors).replace("Yuanyuan Hu", "<b>Yuanyuan Hu</b>")


def pdf_footer(canvas, doc, data):
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#5f666e"))
    footer = f"Latest CV: {norm_url(data['site']['url'])}    Page {doc.page}"
    canvas.drawCentredString(letter[0] / 2, 0.38 * inch, footer)
    canvas.restoreState()


def build_pdf(data):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("Name", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20, leading=22, textColor=colors.HexColor("#1e3a5f"), alignment=TA_CENTER, spaceAfter=2))
    styles.add(ParagraphStyle("Sub", parent=styles["Normal"], fontName="Helvetica-Oblique", fontSize=10.5, leading=13, textColor=colors.HexColor("#3a5e8c"), alignment=TA_CENTER, spaceAfter=1))
    styles.add(ParagraphStyle("Contact", parent=styles["Normal"], fontName="Helvetica", fontSize=8.4, leading=10, textColor=colors.HexColor("#3a5e8c"), alignment=TA_CENTER, spaceAfter=8))
    styles.add(ParagraphStyle("H", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=10, leading=12, textColor=colors.HexColor("#1e3a5f"), spaceBefore=8, spaceAfter=3))
    styles.add(ParagraphStyle("BodyX", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.9, leading=11.2, textColor=colors.HexColor("#262a2e"), spaceAfter=4))
    styles.add(ParagraphStyle("Small", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.6, leading=9.2, textColor=colors.HexColor("#5f666e"), spaceAfter=3))
    styles.add(ParagraphStyle("Pub", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.9, leading=9.4, textColor=colors.HexColor("#262a2e"), leftIndent=14, firstLineIndent=-14, spaceAfter=2.8))
    styles.add(ParagraphStyle("Link", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.2, leading=8.5, textColor=colors.HexColor("#3a5e8c"), leftIndent=14, spaceAfter=2))

    story = [
        para(safe(data["name"]), styles["Name"]),
        para(safe(data["title"]), styles["Sub"]),
        para(safe("  |  ".join(x for x in [data.get("email"), data.get("orcid"), data.get("researchgate"), data.get("location")] if x)), styles["Contact"]),
        para("RESEARCH PROFILE", styles["H"]),
        para(safe(data.get("profile", "")), styles["BodyX"]),
        para("RESEARCH POSITIONING", styles["H"]),
        para("<b>Recommended primary positioning:</b> " + safe(data["positioningOptions"][0]), styles["BodyX"]),
    ]
    for option in data.get("positioningOptions", [])[1:]:
        story.append(para("- " + safe(option), styles["Small"]))

    story.extend([para("EDUCATION", styles["H"])])
    for e in data["education"]:
        story.append(para(f"<b>{safe(e['degree'])}</b> &nbsp; <font color='#3a5e8c'>{safe(e['period'])}</font><br/>{safe(e['institution'])}, {safe(e.get('location', ''))}<br/><font color='#5f666e'>{safe(e.get('detail', ''))} Supervisor: {safe(e.get('supervisor', ''))}.</font>", styles["BodyX"]))

    story.append(para("RESEARCH INTERESTS AND METHODS", styles["H"]))
    story.append(para("<b>Interests:</b> " + safe("; ".join(data.get("interests", []))) + ".", styles["BodyX"]))
    story.append(para("<b>Methods:</b> " + safe(" ".join(data.get("skills", []))), styles["BodyX"]))

    story.append(para("AWARDS AND HONOURS", styles["H"]))
    for award in data.get("awards", []):
        story.append(para(f"<font color='#3a5e8c'><b>{safe(award['year'])}</b></font> &nbsp; {safe(award['text'])}", styles["Small"]))

    story.append(para("PUBLICATIONS", styles["H"]))
    pubs = data.get("publications", {})
    for label, key in (("Lead-author publications", "lead"), ("Collaborative publications", "collaborative")):
        if not pubs.get(key):
            continue
        story.append(para(f"<b>{label}</b>", styles["Small"]))
        for i, pub in enumerate(pubs[key], 1):
            block = []
            details = " ".join(f"[{safe(x)}]" for x in [pub.get("contribution"), pub.get("metrics")] if x)
            block.append(para(f"{i}. {highlighted_authors(pub['authors'])} ({safe(pub['year'])}). {safe(pub['title'])}. <i>{safe(pub['venue'])}</i>, {safe(pub.get('info', ''))}. <font color='#5f666e'>{details}</font>", styles["Pub"]))
            link = doi_url(pub)
            if link:
                block.append(para(f"DOI/article: <link href='{safe(link)}'>{safe(link)}</link>", styles["Link"]))
            story.append(KeepTogether(block))

    doc = SimpleDocTemplate(
        PDF_OUT,
        pagesize=letter,
        rightMargin=0.58 * inch,
        leftMargin=0.58 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.65 * inch,
        title="Yuanyuan Hu Extended CV",
        author="Yuanyuan Hu",
    )
    doc.build(story, onFirstPage=lambda c, d: pdf_footer(c, d, data), onLaterPages=lambda c, d: pdf_footer(c, d, data))


def main():
    with open(DATA_PATH, encoding="utf-8") as f:
        data = json.load(f)
    build_docx(data)
    build_pdf(data)
    print(DOCX_OUT)
    print(PDF_OUT)


if __name__ == "__main__":
    main()
