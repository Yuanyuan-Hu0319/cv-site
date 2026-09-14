#!/usr/bin/env python3
"""Build an editable, professional ONE-PAGE CV (.docx) from cv-data.json.

Dense but elegant single-page layout: tight margins, compact short sections,
hanging-indent references, a single equal-contribution footnote.

Run:  python3 build_docx.py
PDF:  soffice --headless --convert-to pdf cv.docx
"""
import json
import os

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
NAVY = RGBColor(0x1E, 0x3A, 0x5F)
SOFT_BLUE = RGBColor(0x3A, 0x5E, 0x8C)
INK = RGBColor(0x26, 0x2A, 0x2E)
MUTED = RGBColor(0x5F, 0x66, 0x6E)
BODY_FONT = "Calibri"
HEAD_FONT = "Cambria"

# Type sizes tuned for a single page.
S_BODY = 9.0
S_PUB = 8.1
S_SMALL = 7.3


def set_margins(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.4)
        s.bottom_margin = Inches(0.35)
        s.left_margin = Inches(0.55)
        s.right_margin = Inches(0.55)


def run(p, text, *, bold=False, italic=False, size=S_BODY, color=INK, font=BODY_FONT):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.name = font
    return r


def tight(p, before=0, after=2, line=1.0):
    if p.alignment is None:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C9D3E0")
    pbdr.append(bottom)
    pPr.append(pbdr)


def section_title(doc, text):
    p = doc.add_paragraph()
    tight(p, before=5, after=2)
    run(p, text.upper(), bold=True, size=9.5, color=NAVY, font=HEAD_FONT)
    bottom_border(p)
    return p


def authors_runs(p, authors, highlight, size=S_PUB):
    if highlight and highlight in authors:
        parts = authors.split(highlight)
        for i, part in enumerate(parts):
            if part:
                run(p, part, size=size, color=INK)
            if i < len(parts) - 1:
                run(p, highlight, bold=True, size=size, color=NAVY)
    else:
        run(p, authors, size=size, color=INK)


def add_pub(doc, pub, idx, highlight):
    p = doc.add_paragraph()
    tight(p, before=0, after=2, line=1.06)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    run(p, f"{idx}. ", bold=True, size=S_PUB, color=SOFT_BLUE)
    authors_runs(p, pub["authors"], highlight)
    run(p, f" ({pub['year']}). ", size=S_PUB, color=INK)
    run(p, pub["title"] + ". ", size=S_PUB, color=INK)
    run(p, pub["venue"], italic=True, size=S_PUB, color=INK)
    if pub.get("info"):
        run(p, f", {pub['info']}", size=S_PUB, color=INK)
    run(p, ". ", size=S_PUB, color=INK)
    if pub.get("contribution"):
        run(p, f"[{pub['contribution']}] ", size=S_SMALL, color=MUTED)
    if pub.get("metrics"):
        run(p, f"[{pub['metrics']}]", size=S_SMALL, color=MUTED)


def main():
    data = json.load(open(os.path.join(HERE, "cv-data.json"), encoding="utf-8"))
    doc = Document()
    set_margins(doc)
    doc.styles["Normal"].font.name = BODY_FONT
    doc.styles["Normal"].font.size = Pt(S_BODY)

    # --- header ---
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; tight(p, after=0)
    run(p, data["name"], bold=True, size=19, color=NAVY, font=HEAD_FONT)
    if data.get("nameZh"):
        run(p, "  " + data["nameZh"], size=12, color=MUTED, font=HEAD_FONT)
    if data.get("title"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; tight(p, after=0)
        run(p, data["title"], italic=True, size=10, color=SOFT_BLUE, font=HEAD_FONT)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; tight(p, after=0)
    run(p, data["affiliation"], size=8.5, color=MUTED)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; tight(p, after=1)
    extras = []
    if data.get("orcid"):
        extras.append(data["orcid"].replace("https://", ""))
    if data.get("researchgate"):
        extras.append(data["researchgate"].replace("https://www.", "").replace("https://", ""))
    line = "  ·  ".join([data.get("email", "")] + extras + [data.get("location", "")])
    run(p, line, size=8.5, color=SOFT_BLUE)

    # --- education ---
    section_title(doc, "Education")
    for e in data["education"]:
        p = doc.add_paragraph(); tight(p, after=0)
        run(p, e["degree"], bold=True, size=S_BODY, color=INK)
        run(p, f"   {e['period']}", size=8.3, color=SOFT_BLUE)
        p2 = doc.add_paragraph(); tight(p2, after=0)
        run(p2, e["institution"], size=8.7, color=INK)
        if e.get("location"):
            run(p2, f" · {e['location']}", size=8.7, color=MUTED)
        if e.get("detail"):
            p3 = doc.add_paragraph(); tight(p3, after=0)
            run(p3, e["detail"], size=8.4, color=MUTED)
        if e.get("supervisor"):
            p4 = doc.add_paragraph(); tight(p4, after=3)
            run(p4, "Supervisor: " + e["supervisor"], size=8.4, color=MUTED)

    # --- research interests (one compact line) ---
    if data.get("interests"):
        section_title(doc, "Research Interests")
        p = doc.add_paragraph(); tight(p, after=2)
        run(p, "  ·  ".join(data["interests"]), size=S_BODY, color=INK)

    # --- skills (compact paragraph, not bullets) ---
    if data.get("skills"):
        section_title(doc, "Skills")
        p = doc.add_paragraph(); tight(p, after=2, line=1.05)
        run(p, "  ·  ".join(s.rstrip(".") for s in data["skills"]) + ".", size=8.6, color=INK)

    # --- languages (one line) ---
    if data.get("languages"):
        section_title(doc, "Languages")
        p = doc.add_paragraph(); tight(p, after=2)
        for i, l in enumerate(data["languages"]):
            if i:
                run(p, "    ", size=S_BODY)
            run(p, l["name"] + ": ", bold=True, size=S_BODY, color=INK)
            run(p, l.get("level", "") + (f" ({l['note']})" if l.get("note") else ""), size=S_BODY, color=MUTED)

    # --- awards ---
    if data.get("awards"):
        section_title(doc, "Awards & Honours")
        for a in data["awards"]:
            p = doc.add_paragraph(); tight(p, after=1, line=1.02)
            p.paragraph_format.left_indent = Inches(0.42)
            p.paragraph_format.first_line_indent = Inches(-0.42)
            run(p, f"{a['year']}  ", bold=True, size=8.5, color=SOFT_BLUE)
            run(p, a["text"], size=8.5, color=INK)

    # --- publications ---
    pubs = data.get("publications", {})
    hl = data.get("highlightName")
    has_equal = any(pp.get("equalContrib") for g in ("lead", "collaborative") for pp in pubs.get(g, []))
    if pubs.get("lead") or pubs.get("collaborative"):
        section_title(doc, "Publications")
        if pubs.get("lead"):
            p = doc.add_paragraph(); tight(p, before=1, after=1)
            run(p, "Lead-author publications", bold=True, size=8.4, color=MUTED)
            for i, pub in enumerate(pubs["lead"], 1):
                add_pub(doc, pub, i, hl)
        if pubs.get("collaborative"):
            p = doc.add_paragraph(); tight(p, before=2, after=1)
            run(p, "Collaborative publications", bold=True, size=8.4, color=MUTED)
            for i, pub in enumerate(pubs["collaborative"], 1):
                add_pub(doc, pub, i, hl)
        if has_equal:
            p = doc.add_paragraph(); tight(p, before=2, after=0)
            run(p, "# indicates equal contribution.", italic=True, size=S_SMALL, color=MUTED)

    if data.get("site", {}).get("url"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; tight(p, before=2, after=0)
        run(p, "Latest CV: " + data["site"]["url"].replace("https://", "").rstrip("/") + "/", size=6.6, color=MUTED)

    out = os.path.join(HERE, "cv.docx")
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
