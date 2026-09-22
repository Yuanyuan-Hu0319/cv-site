#!/usr/bin/env python3
"""Build the one-page and extended CVs from cv-data.json.

The DOCX files are the source artifacts. Export them with Microsoft Word so
hyperlinks, pagination, and font metrics remain consistent in the PDFs.
"""

import json
import os
import shutil

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Mm, Pt, RGBColor


HERE = os.path.dirname(os.path.abspath(__file__))
DATA_PATH = os.path.join(HERE, "cv-data.json")
ONE_PAGE_OUT = os.path.join(HERE, "Yuanyuan_Hu_CV.docx")
EXTENDED_OUT = os.path.join(HERE, "Yuanyuan_Hu_CV_Extended.docx")
LEGACY_OUT = os.path.join(HERE, "cv.docx")

BODY_FONT = "Aptos"
HEAD_FONT = "Aptos Display"
INK = RGBColor(0x20, 0x25, 0x2B)
NAVY = RGBColor(0x19, 0x36, 0x54)
TEAL = RGBColor(0x1E, 0x6A, 0x73)
MUTED = RGBColor(0x5C, 0x64, 0x6D)
PALE = RGBColor(0xD8, 0xE0, 0xE5)


def load_data():
    with open(DATA_PATH, encoding="utf-8") as handle:
        return json.load(handle)


def set_run_font(run, name=BODY_FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def add_run(paragraph, text, *, bold=False, italic=False, size=9.3,
            color=INK, font=BODY_FONT):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    set_run_font(run, font)
    return run


def add_hyperlink(paragraph, text, url, *, size=9.0, bold=False,
                  italic=False, color=TEAL, underline=True):
    relationship = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), BODY_FONT)
    fonts.set(qn("w:hAnsi"), BODY_FONT)
    run_properties.append(fonts)

    color_element = OxmlElement("w:color")
    color_element.set(qn("w:val"), str(color))
    run_properties.append(color_element)
    size_element = OxmlElement("w:sz")
    size_element.set(qn("w:val"), str(int(size * 2)))
    run_properties.append(size_element)
    size_complex = OxmlElement("w:szCs")
    size_complex.set(qn("w:val"), str(int(size * 2)))
    run_properties.append(size_complex)
    if bold:
        run_properties.append(OxmlElement("w:b"))
    if italic:
        run_properties.append(OxmlElement("w:i"))
    if underline:
        underline_element = OxmlElement("w:u")
        underline_element.set(qn("w:val"), "single")
        run_properties.append(underline_element)

    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(run_properties)
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def set_paragraph(paragraph, *, before=0, after=3, line=1.08,
                  alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, keep=False):
    paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_together = keep


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def set_page(section, *, top, bottom, left, right):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.22)


def configure_document(doc, *, top, bottom, left, right):
    set_page(doc.sections[0], top=top, bottom=bottom, left=left, right=right)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.08

    title = doc.styles["Title"]
    title.font.name = HEAD_FONT
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = NAVY
    title.paragraph_format.space_after = Pt(0)

    if "CV Section" not in doc.styles:
        style = doc.styles.add_style("CV Section", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles["CV Section"]
    style.font.name = HEAD_FONT
    style.font.size = Pt(10.2)
    style.font.bold = True
    style.font.color.rgb = NAVY
    style.paragraph_format.space_before = Pt(6)
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.keep_with_next = True

    properties = doc.core_properties
    properties.title = "Yuanyuan Hu Curriculum Vitae"
    properties.author = "Yuanyuan Hu"
    properties.subject = "Academic curriculum vitae"
    properties.keywords = "neuroimaging, twin studies, behavioral genetics, anxiety, depression"


def section_title(doc, text, *, before=None):
    paragraph = doc.add_paragraph(style="CV Section")
    if before is not None:
        paragraph.paragraph_format.space_before = Pt(before)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_run(paragraph, text.upper(), bold=True, size=10.2, color=NAVY, font=HEAD_FONT)
    return paragraph


def clear_paragraph_borders(paragraph):
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is not None:
        properties.remove(borders)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right", "between"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    properties.append(borders)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    run.font.size = Pt(7.2)
    run.font.color.rgb = MUTED
    set_run_font(run)


def add_footer(section, site_url, *, show_page=True):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    add_run(paragraph, "Latest CV: ", size=7.2, color=MUTED)
    label = site_url.replace("https://", "").rstrip("/") + "/"
    add_hyperlink(paragraph, label, site_url, size=7.2, color=TEAL)
    if show_page:
        add_run(paragraph, "   |   Page ", size=7.2, color=MUTED)
        add_page_field(paragraph)


def add_header(doc, data, *, compact=False):
    name = doc.add_paragraph(style="Title")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_before = Pt(0)
    clear_paragraph_borders(name)
    add_run(name, data["name"], bold=True, size=21 if compact else 22,
            color=NAVY, font=HEAD_FONT)

    role = doc.add_paragraph()
    set_paragraph(role, after=1, line=1.0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_run(role, data["title"], bold=True, size=10.0, color=TEAL, font=HEAD_FONT)
    add_run(role, "  |  ", size=9.0, color=PALE)
    add_run(role, data["location"], size=9.0, color=MUTED)

    affiliation = doc.add_paragraph()
    set_paragraph(affiliation, after=1, line=1.0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_run(affiliation, data["affiliation"], size=8.7 if compact else 9.0, color=MUTED)

    contacts = doc.add_paragraph()
    set_paragraph(contacts, after=3 if compact else 5, line=1.0,
                  alignment=WD_ALIGN_PARAGRAPH.CENTER)
    links = [
        ("Email", "mailto:" + data["email"]),
        ("Website", data["site"]["url"]),
        ("ORCID", data["orcid"]),
        ("ResearchGate", data["researchgate"]),
    ]
    for index, (label, url) in enumerate(links):
        if index:
            add_run(contacts, "   |   ", size=8.2, color=PALE)
        add_hyperlink(contacts, label, url, size=8.5, bold=True, color=TEAL)


def add_highlighted_authors(paragraph, authors, highlight, *, size):
    if not highlight or highlight not in authors:
        add_run(paragraph, authors, size=size)
        return
    chunks = authors.split(highlight)
    for index, chunk in enumerate(chunks):
        if chunk:
            add_run(paragraph, chunk, size=size)
        if index < len(chunks) - 1:
            add_run(paragraph, highlight, bold=True, size=size, color=NAVY)


def publication_link(pub):
    links = pub.get("links") or {}
    url = links.get("doi") or links.get("preprint") or ""
    if url and not url.startswith("http"):
        return "https://doi.org/" + url
    return url


def add_publication(doc, pub, number, highlight, *, compact=False):
    size = 8.95 if compact else 8.95
    paragraph = doc.add_paragraph()
    set_paragraph(paragraph, after=3.2 if compact else 4, line=1.1, keep=True)
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.first_line_indent = Inches(-0.22)
    add_run(paragraph, f"{number}. ", bold=True, size=size, color=TEAL)
    author_text = pub.get("shortAuthors", pub["authors"]) if compact else pub["authors"]
    add_highlighted_authors(paragraph, author_text, highlight, size=size)
    add_run(paragraph, f" ({pub['year']}). ", size=size)
    link = publication_link(pub)
    if link:
        add_hyperlink(paragraph, pub["title"], link, size=size, color=TEAL)
    else:
        add_run(paragraph, pub["title"], size=size)
    add_run(paragraph, ". ", size=size)
    add_run(paragraph, pub["venue"], italic=True, size=size)
    if pub.get("info"):
        add_run(paragraph, f", {pub['info']}", size=size)
    add_run(paragraph, ".", size=size)
    if compact:
        details = [pub.get("contribution"), pub.get("metrics")]
        add_run(paragraph, "  " + " | ".join(x for x in details if x),
                size=8.0, color=MUTED)
    else:
        details = [pub.get("contribution"), pub.get("metrics")]
        if any(details):
            add_run(paragraph, "  " + " | ".join(x for x in details if x),
                    size=8.0, color=MUTED)
        if link:
            add_run(paragraph, "  ", size=8.0)
            add_hyperlink(paragraph, "[Article]", link, size=8.0,
                          bold=True, color=TEAL, underline=False)


def add_education(doc, data, *, compact=False):
    section_title(doc, "Education", before=3 if compact else None)
    for item in data["education"]:
        heading = doc.add_paragraph()
        set_paragraph(heading, after=0, line=1.0, keep=True)
        heading.paragraph_format.tab_stops.add_tab_stop(
            Inches(6.9 if compact else 6.75), WD_TAB_ALIGNMENT.RIGHT
        )
        add_run(heading, item["degree"], bold=True,
                size=9.35 if compact else 9.4, color=INK)
        add_run(heading, "\t" + item["period"], bold=True,
                size=8.6 if compact else 8.7, color=TEAL)

        detail = doc.add_paragraph()
        set_paragraph(detail, after=3 if compact else 4,
                      line=1.08 if compact else 1.1, keep=True)
        add_run(detail, item["institution"], italic=True,
                size=8.95 if compact else 9.0, color=MUTED)
        if item.get("location"):
            add_run(detail, f", {item['location']}. ", size=8.95 if compact else 9.0, color=MUTED)
        add_run(detail, item.get("detail", ""), size=8.95 if compact else 9.0)
        if item.get("supervisor"):
            add_run(detail, f" Supervisor: {item['supervisor']}.",
                    size=8.95 if compact else 9.0, color=MUTED)


def build_one_page(data):
    doc = Document()
    configure_document(doc, top=0.38, bottom=0.4, left=0.58, right=0.58)
    add_footer(doc.sections[0], data["site"]["url"], show_page=False)
    add_header(doc, data, compact=True)

    section_title(doc, "Research Profile", before=2)
    profile = doc.add_paragraph()
    set_paragraph(profile, after=3, line=1.12)
    add_run(profile, data["profile"], size=9.25)

    add_education(doc, data, compact=True)

    section_title(doc, "Research Focus and Methods", before=4)
    focus = doc.add_paragraph()
    set_paragraph(focus, after=2, line=1.08)
    add_run(focus, "Focus: ", bold=True, size=9.1, color=NAVY)
    add_run(focus, "; ".join(data["interests"]) + ".", size=9.1)
    methods = doc.add_paragraph()
    set_paragraph(methods, after=3, line=1.08)
    add_run(methods, "Methods: ", bold=True, size=9.1, color=NAVY)
    add_run(methods, " ".join(data["skills"]), size=9.1)

    section_title(doc, "Selected Awards and Honours", before=4)
    for award in data["awards"]:
        paragraph = doc.add_paragraph()
        set_paragraph(paragraph, after=1, line=1.02, keep=True)
        paragraph.paragraph_format.left_indent = Inches(0.43)
        paragraph.paragraph_format.first_line_indent = Inches(-0.43)
        add_run(paragraph, f"{award['year']}  ", bold=True, size=8.6, color=TEAL)
        add_run(paragraph, award["text"], size=8.6)

    language = doc.add_paragraph()
    set_paragraph(language, after=2, line=1.0)
    add_run(language, "Languages: ", bold=True, size=8.6, color=NAVY)
    add_run(language, "English - Advanced (IELTS 7.0).", size=8.6)

    section_title(doc, "Selected Publications", before=4)
    for index, pub in enumerate(data["publications"]["lead"], 1):
        add_publication(doc, pub, index, data.get("highlightName"), compact=True)

    note = doc.add_paragraph()
    set_paragraph(note, before=1, after=0, line=1.0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_run(note, "# Equal contribution. Full author lists and collaborative publications: ",
            italic=True, size=7.3, color=MUTED)
    add_hyperlink(note, "Extended CV", data["site"]["url"] + data["site"]["extendedPdf"],
                  size=7.3, color=TEAL)
    add_run(note, "  |  ", size=7.3, color=PALE)
    add_hyperlink(note, "ORCID", data["orcid"], size=7.3, color=TEAL)

    doc.save(ONE_PAGE_OUT)
    shutil.copy2(ONE_PAGE_OUT, LEGACY_OUT)


def build_extended(data):
    doc = Document()
    configure_document(doc, top=0.56, bottom=0.58, left=0.68, right=0.68)
    add_footer(doc.sections[0], data["site"]["url"], show_page=True)
    add_header(doc, data, compact=False)

    section_title(doc, "Research Profile")
    profile = doc.add_paragraph()
    set_paragraph(profile, after=4, line=1.13)
    add_run(profile, data["profile"], size=9.35)

    add_education(doc, data, compact=False)

    section_title(doc, "Research Focus")
    focus = doc.add_paragraph()
    set_paragraph(focus, after=4, line=1.12)
    add_run(focus, "; ".join(data["interests"]) + ".", size=9.25)

    section_title(doc, "Methods and Analytical Tools")
    methods = doc.add_paragraph()
    set_paragraph(methods, after=4, line=1.12)
    add_run(methods, " ".join(data["skills"]), size=9.25)

    section_title(doc, "Awards and Honours")
    for award in data["awards"]:
        paragraph = doc.add_paragraph()
        set_paragraph(paragraph, after=2, line=1.08, keep=True)
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.5)
        add_run(paragraph, f"{award['year']}  ", bold=True, size=8.75, color=TEAL)
        add_run(paragraph, award["text"], size=8.75)

    languages = doc.add_paragraph()
    set_paragraph(languages, before=3, after=0, line=1.05)
    add_run(languages, "Languages: ", bold=True, size=8.8, color=NAVY)
    language_text = "; ".join(
        f"{item['name']} - {item.get('level', '')}"
        + (f" ({item['note']})" if item.get("note") else "")
        for item in data.get("languages", [])
    )
    add_run(languages, language_text + ".", size=8.8)

    page_break = doc.add_paragraph()
    page_break.paragraph_format.space_after = Pt(0)
    page_break.add_run().add_break(WD_BREAK.PAGE)

    section_title(doc, "Publications", before=0)
    intro = doc.add_paragraph()
    set_paragraph(intro, after=4, line=1.0)
    add_run(intro, "Yuanyuan Hu is shown in bold; # indicates equal contribution. Publication titles link to the article or preprint.",
            italic=True, size=8.2, color=MUTED)

    publications = data["publications"]
    for label, key in (("Lead-author publications", "lead"),
                       ("Collaborative publications", "collaborative")):
        subheading = doc.add_paragraph()
        set_paragraph(subheading, before=3, after=2, line=1.0,
                      alignment=WD_ALIGN_PARAGRAPH.LEFT)
        keep_with_next(subheading)
        add_run(subheading, label, bold=True, size=9.3, color=NAVY, font=HEAD_FONT)
        for index, pub in enumerate(publications[key], 1):
            add_publication(doc, pub, index, data.get("highlightName"), compact=False)

    doc.save(EXTENDED_OUT)


def main():
    data = load_data()
    build_one_page(data)
    build_extended(data)
    print(ONE_PAGE_OUT)
    print(EXTENDED_OUT)


if __name__ == "__main__":
    main()
